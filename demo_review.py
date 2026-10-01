"""SIMULATED reviewer for the sample data only: fills the QA report's Decision column using the sample answer key
(data/_truth_<month>.csv), the way a person checking the platforms would. Never use this with real data.

    python demo_review.py --month 2026-09            # fills a copy of the QA report and resumes the run
"""
import argparse
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from pipeline import ROOT
from review import resume


def fill(month: str, partial: float = 1.0) -> Path:
    src = ROOT / "output" / month / f"qa_report_{month}.xlsx"
    dst = src.with_name(f"qa_report_{month}_reviewed.xlsx")
    shutil.copy(src, dst)
    t = pd.read_csv(ROOT / "data" / f"_truth_{month}.csv").set_index(["creator_id", "platform"])
    t["s"] = t.saves_90d.fillna(0)
    truth = {"follower_count": t.follower_count, "reach_rate": t.views_90d / t.follower_count,
             "growth_rate": t.net_new_followers_90d / (t.follower_count - t.net_new_followers_90d),
             "meaningful_eng_rate": (t.comments_90d + t.shares_90d + t.s) / (t.likes_90d + t.comments_90d + t.shares_90d + t.s)}
    wb = load_workbook(dst)
    ws = wb["Blocked — review"]
    hdr = {c.value: i + 1 for i, c in enumerate(ws[4])}
    vendors = [h[:-6] for h in hdr if str(h).endswith(" value")]
    rows = [r for r in range(5, ws.max_row + 1) if ws.cell(r, hdr["Status"]).value == "open"]
    for r in rows[: int(len(rows) * partial)]:
        cid, plat = ws.cell(r, hdr["Creator ID"]).value, ws.cell(r, hdr["Channel"]).value
        code = ws.cell(r, hdr["Reason code"]).value
        vals = {v: ws.cell(r, hdr[f"{v} value"]).value for v in vendors if ws.cell(r, hdr[f"{v} value"]).value is not None}
        if code == "IDENTITY_MATCH_UNCERTAIN":
            real = t.loc[(cid, plat), "handle"].lower() if (cid, plat) in t.index else None
            match = [v for v, h in vals.items() if str(h).lower() == (real or "")]
            dec, man = (f"SELECT_{match[0].upper()}", None) if match else ("EXCLUDE", None)
        else:
            fld = ws.cell(r, hdr["Field(s) held back"]).value.split(",")[0].strip()
            tv = truth[fld].get((cid, plat))
            num = {v: float(str(x).replace(",", "")) for v, x in vals.items()}
            best = min(num, key=lambda v: abs(num[v] - tv) / max(abs(tv), 1e-9))
            close = abs(num[best] - tv) / max(abs(tv), 1e-9) < 0.10
            dec, man = (f"SELECT_{best.upper()}", None) if close else ("MANUAL_VALUE", round(float(tv), 4))
        ws.cell(r, hdr["Decision"]).value = dec
        ws.cell(r, hdr["Manual value"]).value = man
        ws.cell(r, hdr["Reviewer"]).value = "Simulated reviewer"
        ws.cell(r, hdr["Note"]).value = "checked against the platform (sample answer key)"
    wb.save(dst)
    return dst


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", required=True)
    ap.add_argument("--partial", type=float, default=1.0, help="share of open items to decide (to test partial review)")
    a = ap.parse_args()
    resume(a.month, fill(a.month, a.partial))
