"""Generate Demo Data: a complete SYNTHETIC monthly drop for demonstrating the prototype.

Creates, in demo_workspace/ (kept apart from the real memory/):
  inputs/<month>/  SYNTHETIC_DEMO_creator_request.xlsx + SYNTHETIC_DEMO vendor A / B / C files (200–300 rows each)
  memory/          3 months of synthetic history: vendor QA scores, reviewed disputes (review audit), creator scores
  scenarios_<month>.json   which creators were seeded for which QA scenario (for the presenter)

Names, IDs and numbers are random on every generation, but these QA scenarios are ALWAYS seeded:
  1 clean agreement (GREEN)            2 Tier 1 missing a field → fallback (YELLOW PRIMARY_SOURCE_MISSING)
  3 material discrepancy (RED)         + a Tier 1 outlier overruled by two agreeing vendors (YELLOW)
  4 identity can't be matched (RED)    + an identity conflict resolved by majority (YELLOW)
  5 historical anomaly (YELLOW)        6 different vendor quality: A reliable, B middling, C weak (3-month history)
Everything is synthetic: invented names, "demo" handles, DEMO- IDs. No real creators, vendors or company data.

    python demo_data.py [--month 2026-10] [--seed 7]
"""
from __future__ import annotations

import argparse
import json
import random
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.styles import Font

from request_template import write_request

ROOT = Path(__file__).parent
WORKSPACE = ROOT / "demo_workspace"
LABEL = "SYNTHETIC DEMO DATA"
CFG = yaml.safe_load((ROOT / "config.yaml").read_text())
PLATFORMS = CFG["platforms"]
REACH = {"tiktok": 3.0, "youtube": 1.2, "instagram": 0.6}
COUNTRIES = {"US": "United States", "GB": "United Kingdom", "CA": "Canada", "BR": "Brazil", "IN": "India",
             "DE": "Germany", "FR": "France", "JP": "Japan", "MX": "Mexico", "KR": "South Korea"}
CATS = ["gaming", "beauty", "fitness", "food", "tech", "travel", "music", "education", "comedy", "finance"]
ADJ = ["amber", "velvet", "lunar", "copper", "misty", "neon", "quiet", "solar", "crimson", "frosty", "golden", "hollow",
       "ivory", "jade", "kinetic", "lucky", "mellow", "nimble", "opal", "pixel", "rustic", "silver", "tidal", "urban",
       "vivid", "wild", "zesty", "brisk", "cosmic", "dusky"]
NOUN = ["otter", "comet", "fern", "harbor", "lantern", "maple", "nova", "orchid", "pebble", "quartz", "raven", "sparrow",
        "thistle", "willow", "yarrow", "badger", "cinder", "delta", "ember", "falcon", "glacier", "heron", "iris", "juniper",
        "kestrel", "lotus", "meadow", "nectar", "onyx", "prairie"]

# how many examples of each scenario every generation gets
N_REQUESTED, N_CLEAN, N_FALLBACK, N_DISCREPANCY, N_T1_OUTLIER, N_IDENTITY, N_IDENTITY_MAJ, N_ANOMALY = 100, 8, 6, 4, 2, 3, 2, 4


def _months_before(month: str, n: int) -> list:
    p = pd.Period(month, "M")
    return [str(p - k) for k in range(n, 0, -1)]


def _label_sheet(path: Path, what: str, generated: str):
    wb = load_workbook(path)
    ws = wb.create_sheet("About")
    lines = [LABEL, f"{what}", f"Generated {generated} by the prototype's Generate Demo Data feature.",
             "All creators, vendor records, historical records and metrics are synthetic.",
             "Not real creators, vendors, clients or company data."]
    for i, t in enumerate(lines, start=1):
        ws.cell(i, 1, t).font = Font(name="Arial", bold=i == 1, size=14 if i == 1 else 11, color="C2410C" if i == 1 else None)
    ws.column_dimensions["A"].width = 90
    wb.save(path)


def generate(month: str | None = None, seed: int | None = None, workspace: Path = WORKSPACE) -> dict:
    month = month or date.today().strftime("%Y-%m")
    seed = seed if seed is not None else random.SystemRandom().randrange(1_000_000)
    rng, npr = random.Random(seed), np.random.default_rng(seed)
    end = pd.Period(month, "M").end_time.date()
    generated = datetime.now().strftime("%Y-%m-%d %H:%M")
    if workspace.exists():
        shutil.rmtree(workspace)
    inbox = workspace / "inputs" / month
    mem = workspace / "memory"
    inbox.mkdir(parents=True)
    mem.mkdir(parents=True)

    # ------------------------------------------------------------------ synthetic creators (random every time)
    n_all = 160
    ids = [f"DEMO-{i:04d}" for i in rng.sample(range(1, 10000), n_all)]
    combos = rng.sample([(a, b) for a in ADJ for b in NOUN], n_all)
    creators = pd.DataFrame({
        "creator_id": ids,
        "display_name": [f"{a.title()} {b.title()} (demo)" for a, b in combos],
        "slug": [f"demo_{a}{b}" for a, b in combos],
        "country": [rng.choice(list(COUNTRIES)) for _ in ids],
        "contact_email": [f"{a}.{b}@example.invalid" for a, b in combos],
        "category": [rng.choice(CATS) for _ in ids],
    })
    rows = []
    for c in creators.itertuples(index=False):
        for pl in rng.sample(PLATFORMS, k=rng.choices([1, 2, 3], weights=[2, 3, 5])[0]):
            f = int(npr.lognormal(10.8, 1.3))
            g = float(np.clip(npr.normal(0.06, 0.05) + 0.5 / np.log10(max(f, 100)) ** 2, -0.2, 1.5))
            views = int(f * npr.lognormal(np.log(REACH[pl]), 0.6))
            eng = int(views * npr.uniform(0.02, 0.08))
            mix = npr.dirichlet([16, 1.2, 0.8, 1.0 if pl != "youtube" else 1e-4])
            likes, comments, shares, saves = (int(eng * m) for m in mix)
            rows.append({"creator_id": c.creator_id, "platform": pl, "handle": f"@{c.slug}_{pl}", "follower_count": f,
                         "net_new_followers_90d": f - int(round(f / (1 + g))), "views_90d": views, "likes_90d": likes,
                         "comments_90d": comments, "shares_90d": shares,
                         "saves_90d": saves if pl != "youtube" else None,
                         "likes_hidden": pl == "instagram" and rng.random() < 0.08})
    truth = pd.DataFrame(rows).merge(creators.drop(columns="slug"), on="creator_id").set_index(["creator_id", "platform"])
    have = truth.reset_index().groupby("creator_id")["platform"].apply(list).to_dict()

    # ------------------------------------------------------------------ request: 60% need all three channels
    three = [c for c in ids if len(have[c]) == 3]
    full = rng.sample(three, int(N_REQUESTED * 0.6))
    rest = [c for c in ids if c not in full]
    two = rng.sample([c for c in rest if len(have[c]) >= 2], int(N_REQUESTED * 0.2))
    one = rng.sample([c for c in rest if c not in two], N_REQUESTED - len(full) - len(two))
    req = {c: list(PLATFORMS) for c in full} | {c: rng.sample(have[c], 2) for c in two} | {c: rng.sample(have[c], 1) for c in one}
    req_pairs = [(c, pl) for c, pls in req.items() for pl in pls]

    # ------------------------------------------------------------------ pick the seeded scenario cases
    pool = [p for p in req_pairs if not truth.at[p, "likes_hidden"]]
    rng.shuffle(pool)
    used_creators = set()

    def take(n, need=lambda p: True, whole_creator=False):
        out = []
        for p in pool:
            if len(out) >= n:
                break
            if p[0] in used_creators or not need(p):
                continue
            out.append(p)
            used_creators.add(p[0])
        return out

    clean_creators = []
    for c in rng.sample(full, len(full)):                       # whole creators, every channel clean
        if len(clean_creators) >= N_CLEAN:
            break
        if not any(truth.at[(c, pl), "likes_hidden"] for pl in req[c]):
            clean_creators.append(c)
            used_creators.add(c)
    scen = {
        "clean": [(c, pl) for c in clean_creators for pl in req[c]],
        "fallback": take(N_FALLBACK),
        "discrepancy": take(N_DISCREPANCY),
        "t1_outlier": take(N_T1_OUTLIER),
        "identity": take(N_IDENTITY),
        "identity_majority": take(N_IDENTITY_MAJ),
        "anomaly": take(N_ANOMALY, lambda p: truth.at[p, "follower_count"] > 5000),
    }
    seeded = {p for v in scen.values() for p in v}

    # ------------------------------------------------------------------ which vendor has which channel
    # natural coverage (quality tiers by design: A best, B middle, C weakest)
    cover = {"vendor_a": 0.85, "vendor_b": 0.75, "vendor_c": 0.65}
    stale = {"vendor_a": 0.04, "vendor_b": 0.15, "vendor_c": 0.35}
    nulls = {"vendor_a": 0.01, "vendor_b": 0.04, "vendor_c": 0.10}
    presence = {v: {p for p in req_pairs if p not in seeded and rng.random() < cover[v]} for v in cover}
    extra = [p for p in truth.index if p not in set(req_pairs)]                  # unrequested rows, as vendors send
    for v, n_extra in (("vendor_a", 45), ("vendor_b", 70), ("vendor_c", 85)):
        presence[v] |= set(rng.sample(extra, min(n_extra, len(extra))))
    must = {"vendor_a": set(), "vendor_b": set(), "vendor_c": set()}            # seeded rows: fresh, complete
    for p in scen["clean"] + scen["t1_outlier"] + scen["identity_majority"]:
        for v in must:
            must[v].add(p)
    for p in scen["fallback"] + scen["discrepancy"] + scen["anomaly"]:
        must["vendor_a"].add(p); must["vendor_b"].add(p)
    for p in scen["identity"]:
        must["vendor_a"].add(p); must["vendor_c"].add(p)
    for v in presence:
        presence[v] |= must[v]
        for p in seeded:                                     # seeded channels: only the vendors the scenario needs
            if p not in must[v]:
                presence[v].discard(p)

    def stamp(fresh):
        age = rng.randint(3, 75) if fresh else rng.randint(120, 380)
        return (end - timedelta(days=age)).isoformat()

    def vendor_frame(v):
        recs = []
        for p in sorted(presence[v]):
            r = truth.loc[p].to_dict()
            r["creator_id"], r["platform"] = p
            for col in ["follower_count", "views_90d", "likes_90d", "comments_90d", "shares_90d", "saves_90d"]:
                if r[col] is not None and not pd.isna(r[col]):
                    r[col] = int(round(r[col] * npr.normal(1, 0.012)))            # small methodology differences
            is_seeded = p in must[v]
            if not is_seeded:
                for col in ["follower_count", "views_90d", "likes_90d", "comments_90d", "contact_email", "category"]:
                    if rng.random() < nulls[v]:
                        r[col] = None
            if r["likes_hidden"]:
                r["likes_90d"] = "hidden"
            r["ts"] = stamp(fresh=is_seeded or rng.random() > stale[v])
            recs.append(r)
        return pd.DataFrame(recs).set_index(["creator_id", "platform"], drop=False)

    frames = {v: vendor_frame(v) for v in cover}
    A, B, C = frames["vendor_a"], frames["vendor_b"], frames["vendor_c"]
    manifest = {"label": LABEL, "month": month, "seed": seed, "generated": generated, "scenarios": {}}

    def note(key, p, text):
        manifest["scenarios"].setdefault(key, []).append(
            {"creator_id": p[0], "platform": p[1], "creator": truth.at[p, "display_name"], "what": text})

    for p in scen["clean"]:
        note("1_clean_agreement_GREEN", p, "A, B and C all report this channel within ~2% of each other")
    for p in scen["fallback"]:                               # Tier 1 lacks likes → Hub metric comes from Tier 2
        A.at[p, "likes_90d"] = None
        note("2_primary_source_missing_YELLOW", p, "Vendor A (Tier 1) has no likes count; Vendor B fills the engagement metric")
    for p in scen["discrepancy"]:                            # only A and B, follower counts far apart
        k = rng.choice([rng.uniform(0.45, 0.62), rng.uniform(1.55, 1.9)])
        B.at[p, "follower_count"] = int(A.at[p, "follower_count"] * k)
        note("3_cross_vendor_discrepancy_RED", p, f"A {A.at[p, 'follower_count']:,} vs B {B.at[p, 'follower_count']:,} followers; no third source")
    for p in scen["t1_outlier"]:                             # A disagrees with B and C, which agree
        A.at[p, "follower_count"] = int(A.at[p, "follower_count"] * rng.uniform(0.5, 0.6))
        note("3b_tier1_outlier_overruled_YELLOW", p, "Vendor A (Tier 1) is the outlier; B and C agree, so their value is used")
    others = [p for p in truth.index if p not in seeded]
    for i, p in enumerate(scen["identity"]):                 # A and C point at different accounts
        o = rng.choice([q for q in others if q[1] == p[1]])
        for col in ["handle", "follower_count", "net_new_followers_90d", "views_90d", "likes_90d", "comments_90d",
                    "shares_90d", "saves_90d"]:
            C.at[p, col] = truth.at[o, col]
        note("4_identity_match_uncertain_RED", p, f"A says {A.at[p, 'handle']}, C says {C.at[p, 'handle']}; no majority")
    for p in scen["identity_majority"]:
        o = rng.choice([q for q in others if q[1] == p[1]])
        C.at[p, "handle"] = truth.at[o, "handle"]
        note("4b_identity_resolved_by_majority_YELLOW", p, f"A and B agree on {A.at[p, 'handle']}; C reports {C.at[p, 'handle']}")
    for p in scen["anomaly"]:
        note("5_historical_anomaly_YELLOW", p, "follower count ~5–8x last month's (seeded history); vendors agree, so it's kept")

    # ------------------------------------------------------------------ write vendor files in each vendor's own format
    def write_vendor(v, df, filename):
        cmap = CFG["vendors"][v]["column_map"]
        inv = {field: col for col, field in cmap.items()}
        out = pd.DataFrame(index=df.index)
        d = df.copy()
        if v == "vendor_a":
            d["platform"] = d["platform"].map({"youtube": "YouTube", "instagram": "Instagram", "tiktok": "TikTok"})
        if v == "vendor_b":
            d["platform"] = d["platform"].map({"instagram": "IG", "tiktok": "TT", "youtube": "YT"})
            d["country"] = d["country"].map(COUNTRIES)
            prev = [None if (pd.isna(f) or pd.isna(n)) else f - n for f, n in zip(d["follower_count"], d["net_new_followers_90d"])]
            d["followers_prev"] = [None if x is None else round(x / 1000, 2) for x in prev]
            d["follower_count"] = [None if pd.isna(x) else round(x / 1000, 2) for x in d["follower_count"]]
        if v == "vendor_c":
            d["platform"] = [rng.choice([x, x.upper(), x.capitalize()]) for x in d["platform"]]
            d["country"] = d["country"].str.lower()
            d["display_name"] = [n.lower() if rng.random() < .3 else n for n in d["display_name"]]
        for col, field in cmap.items():
            out[col] = d[field].values if field in d else None
        out[CFG["vendors"][v]["last_updated_col"]] = d["ts"].values
        out = out.sample(frac=1, random_state=seed % 1000)
        path = inbox / filename
        out.to_excel(path, index=False, sheet_name="data")
        _label_sheet(path, f"Vendor dataset ({v.replace('_', ' ').title()}) — {len(out)} records for {month}", generated)
        return path

    files = [write_vendor("vendor_a", A, f"SYNTHETIC_DEMO_vendor_A_{month}.xlsx"),
             write_vendor("vendor_b", B, f"SYNTHETIC_DEMO_vendor_B_{month}.xlsx"),
             write_vendor("vendor_c", C, f"SYNTHETIC_DEMO_vendor_C_{month}.xlsx")]
    req_rows = [{"creator_id": c, "creator_name": truth.at[(c, pls[0]), "display_name"], "channels": pls} for c, pls in req.items()]
    rng.shuffle(req_rows)
    rpath = inbox / f"SYNTHETIC_DEMO_creator_request_{month}.xlsx"
    write_request(rpath, PLATFORMS, req_rows, blank_rows=0)
    _label_sheet(rpath, f"Requested Creator List — {len(req_rows)} creators, {len(req_pairs)} channels", generated)
    files.insert(0, rpath)

    # ------------------------------------------------------------------ 3 months of synthetic history in demo memory
    hist_months = _months_before(month, 3)
    quality = {"vendor_a": (0.86, 0.9), "vendor_b": (0.73, 0.6), "vendor_c": (0.57, 0.25)}   # (QA score, review accuracy)
    vrows, audit = [], []
    for k, hm in enumerate(hist_months):
        for t, (v, (q, acc)) in enumerate(quality.items(), start=1):
            score = round(q + npr.normal(0, 0.012), 4)
            vrows.append({"month": hm, "vendor": v, "current_score": score, "blended_score": score, "tier": t,
                          "usable_coverage": round(q - 0.1 + npr.normal(0, .02), 4), "completeness": round(min(.99, q + .08), 4),
                          "freshness": round(min(.99, q + .05), 4), "validity": round(min(.995, q + .1), 4),
                          "implausible_values": int(npr.integers(0, 4 + 10 * t)), "rows": int(npr.integers(200, 300))})
        for j in range(4):                                   # reviewed disputes that month (synthetic evidence)
            p = rng.choice(req_pairs)
            vals = {v: int(truth.at[p, "follower_count"] * npr.normal(1, .02)) for v in quality}
            winner = max(quality, key=lambda v: quality[v][1] * rng.random() ** 0.35)
            audit.append({"month": hm, "reviewed_at": f"{hm}-2{j} 15:0{j}:00", "review_id": f"H{k}{j:03d}",
                          "creator_id": p[0], "platform": p[1], "fields": "follower_count", "code": "CROSS_VENDOR_DISCREPANCY",
                          "original_values": json.dumps(vals), "decision": f"SELECT_{winner.upper()}", "manual_value": None,
                          "final_value": vals[winner], "final_source": winner, "reviewer": "Synthetic history",
                          "note": LABEL})
    pd.DataFrame(vrows).to_csv(mem / "vendor_history.csv", index=False)
    pd.DataFrame(audit).to_csv(mem / "review_audit.csv", index=False)

    crow = []
    pct_base = {p: {m: float(npr.uniform(10, 95)) for m in CFG["metrics"]} for p in truth.index}
    for k, hm in enumerate(hist_months):
        back = len(hist_months) - k                          # months before the demo month
        for c, pls in req.items():
            for pl in have[c]:
                p = (c, pl)
                f_now = truth.at[p, "follower_count"]
                f = f_now / (rng.uniform(5, 8) * (1 + 0.02 * (back - 1))) if p in scen["anomaly"] else f_now / (1.015 ** back)
                row = {"month": hm, "creator_id": c, "platform": pl, "follower_count": int(f),
                       "size_band": None, "complete": True}
                for m in CFG["metrics"]:
                    row[f"{m}__pct"] = round(float(np.clip(pct_base[p][m] + npr.normal(0, 8), 1, 100)), 1)
                row["reach_rate"] = truth.at[p, "views_90d"] / f_now * npr.normal(1, .1)
                row["growth_rate"] = float(np.clip(npr.normal(.06, .04), -.2, 1))
                row["meaningful_eng_rate"] = float(np.clip(npr.normal(.15, .05), .01, .9))
                crow.append(row)
            ch_rows = [r for r in crow if r["month"] == hm and r["creator_id"] == c and r["platform"] in req[c]]
            allr = {"month": hm, "creator_id": c, "platform": "all", "complete": True}
            for m in CFG["metrics"]:
                allr[f"{m}__pct"] = round(float(np.mean([r[f"{m}__pct"] for r in ch_rows])), 1)
            allr["primary_role"] = CFG["metrics"][max(CFG["metrics"], key=lambda m: allr[f"{m}__pct"])]["pillar"]
            crow.append(allr)
    pd.DataFrame(crow).to_csv(mem / "creator_history.csv", index=False)
    manifest["history_months"] = hist_months
    manifest["vendor_quality_history"] = {v: {"qa_score": q, "review_accuracy": a} for v, (q, a) in quality.items()}
    manifest["files"] = [f.name for f in files]
    (workspace / f"scenarios_{month}.json").write_text(json.dumps(manifest, indent=2, default=str))
    truth.reset_index().to_csv(workspace / f"_answer_key_{month}.csv", index=False)   # for verifying the demo only
    return {"files": files, "memory_dir": mem, "out_dir": workspace / "output" / month, "month": month,
            "manifest": manifest, "workspace": workspace}


def reset(workspace: Path = WORKSPACE):
    """Reset Demo: remove generated inputs, demo memory, QA results, review decisions and outputs."""
    if workspace.exists():
        shutil.rmtree(workspace)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--month")
    ap.add_argument("--seed", type=int)
    ap.add_argument("--reset", action="store_true")
    a = ap.parse_args()
    if a.reset:
        reset(); print("Demo reset.")
    else:
        out = generate(a.month, a.seed)
        print(f"{LABEL} for {out['month']} (seed {out['manifest']['seed']}):")
        for f in out["files"]:
            print("  ", f.relative_to(ROOT))
        for k, v in out["manifest"]["scenarios"].items():
            print(f"   {k}: {len(v)}")
