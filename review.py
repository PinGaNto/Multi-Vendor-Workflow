"""Human review of BLOCKED (RED) items: QA report + review form, applying decisions, audit trail, resume.

    python review.py --month 2026-09 --decisions output/2026-09/qa_report_2026-09.xlsx [--reviewer "Name"]

The reviewer fills the 'Blocked — review' tab of the QA report (Decision, Manual value, Reviewer, Note) and
hands the same file back. Only the affected creator-channels are reprocessed. When no blocked item is left open,
the run finalizes: master dataset → percentiles/scores → leaderboard.
"""
from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
BLACK = PatternFill("solid", start_color="0A0A0A")
ORANGE = PatternFill("solid", start_color="EA580C")
INPUT = PatternFill("solid", start_color="FFF1E6")
SEV_FILL = {"RED": PatternFill("solid", start_color="EA580C"), "YELLOW": PatternFill("solid", start_color="FED7AA"),
            "GREEN": PatternFill("solid", start_color="F2F2F2")}
REVIEW_COLS = ["Review ID", "Status", "Creator ID", "Creator", "Channel", "Field(s) held back", "Reason code",
               "Severity"]              # + vendor value columns + the rest
TAIL_COLS = ["Discrepancy", "Source information", "Explanation", "Recommended action",
             "Decision", "Manual value", "Reviewer", "Note", "Final value", "Final source", "Reviewed at"]


def _head(ws, row, headers, widths, fills=None):
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row, c, h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = (fills or {}).get(c, BLACK)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[L(c)].width = w
    ws.row_dimensions[row].height = 32


def _title(ws, t, sub=None):
    ws["A1"] = t
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, italic=True, color="595959")


def _fmt(v):
    if isinstance(v, float):
        return f"{v:,.0f}" if abs(v) >= 100 else f"{v:.4g}"
    return v


def write_qa_sheets(wb, p, review_form: bool = True):
    """Adds 'Blocked — review', 'Warnings' and 'Reason codes' sheets to a workbook."""
    vendors = list(p.cfg["vendors"])
    tier_of = {n: i for i, n in enumerate(p.tiers, start=1)}
    names = p.out_cr["display_name"].to_dict() if hasattr(p, "out_cr") else {}
    flags = p.flags.flags

    # ---- Blocked
    wb_ = wb.create_sheet("Blocked — review")
    _title(wb_, "Blocked records — human review",
           "Only these creator-channels / metrics are held back; everything else was processed. Fill the orange "
           "Decision column (and Manual value where needed), then run review.py with this file.")
    hdr = REVIEW_COLS + [f"{v} value" for v in vendors] + TAIL_COLS
    widths = [9, 10, 11, 16, 10, 22, 26, 9] + [14] * len(vendors) + [11, 26, 60, 44, 20, 12, 14, 24, 14, 18, 18]
    dcol = len(REVIEW_COLS) + len(vendors) + 5
    _head(wb_, 4, hdr, widths, {dcol: ORANGE, dcol + 1: ORANGE, dcol + 2: ORANGE, dcol + 3: ORANGE})
    reds = [f for f in flags if f.severity == "RED"]
    allowed = p.cfg["review"]["decisions"]
    for i, f in enumerate(reds, start=5):
        tier_txt = " > ".join(f"T{tier_of[v]} {v}" for v in p.tiers)
        vals = [_fmt(f.values.get(v)) if f.code != "IDENTITY_MATCH_UNCERTAIN" else
                (("@" + f.values[v]) if v in f.values else None) for v in vendors]
        row = [f.id, f.status, f.creator_id, names.get(f.creator_id), f.platform,
               ", ".join([f.field, *f.dependents]) if f.field != "identity" else "whole channel (identity)",
               f.code, f.severity, *vals,
               (f"{f.discrepancy:.0%}" if f.discrepancy is not None and f.field != "growth_rate" else
                (f"{f.discrepancy * 100:.1f} pts" if f.discrepancy is not None else None)),
               f"Tier order {tier_txt}. Would have used: {f.source_used or '—'}",
               f.explanation, f.action, f.decision or None, f.manual_value, f.reviewer or None, f.note or None,
               _fmt(f.final_value) if f.final_value is not None else None, f.final_source or None, f.reviewed_at or None]
        for c, v in enumerate(row, start=1):
            cell = wb_.cell(i, c, v)
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(wrap_text=c in (len(REVIEW_COLS) + len(vendors) + 3, len(REVIEW_COLS) + len(vendors) + 4),
                                       vertical="top")
        wb_.cell(i, 8).fill = SEV_FILL[f.severity]
        wb_.cell(i, 8).font = Font(name=FONT, bold=True, color="FFFFFF")
        if review_form and f.status == "open":
            for c in range(dcol, dcol + 4):
                wb_.cell(i, c).fill = INPUT
            dv = DataValidation(type="list", formula1='"' + ",".join(allowed.get(f.code, [])) + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Pick a decision",
                                error="Choose one of the listed decisions for this reason code.")
            wb_.add_data_validation(dv)
            dv.add(wb_.cell(i, dcol))
    if not reds:
        wb_.cell(5, 1, "No blocked records this month.").font = Font(name=FONT, italic=True)
    wb_.freeze_panes = "D5"
    wb_.auto_filter.ref = f"A4:{L(len(hdr))}{max(4 + len(reds), 5)}"

    # ---- Warnings
    ww = wb.create_sheet("Warnings")
    _title(ww, "Warnings — recorded, processed automatically",
           "No action needed. Kept for visibility and to monitor vendor quality.")
    hdr = ["Flag ID", "Creator ID", "Creator", "Channel", "Reason code", "Field(s)", "Relevant values",
           "Source selected", "Explanation"]
    _head(ww, 4, hdr, [9, 11, 16, 10, 26, 26, 46, 18, 80])
    yellows = sorted([f for f in flags if f.severity == "YELLOW"], key=lambda f: (f.code, f.creator_id))
    for i, f in enumerate(yellows, start=5):
        rel = "; ".join(f"{k}: {_fmt(v)}" for k, v in f.values.items())
        if f.code == "HISTORICAL_ANOMALY" and f.discrepancy is not None:
            rel += f"; change {f.discrepancy:+.0%}"
        for c, v in enumerate([f.id, f.creator_id, names.get(f.creator_id), f.platform, f.code, f.field, rel,
                               f.source_used, f.explanation], start=1):
            ww.cell(i, c, v).font = Font(name=FONT)
    ww.freeze_panes = "C5"
    ww.auto_filter.ref = f"A4:I{max(4 + len(yellows), 5)}"

    # ---- Reason codes
    wr = wb.create_sheet("Reason codes")
    _title(wr, "Reason codes and severities", "Severities and thresholds are set in config.yaml → qa_rules.")
    _head(wr, 4, ["Reason code", "Severity", "Count this month", "What it means", "What to do"], [28, 10, 12, 70, 70])
    from qa_checks import ACTIONS
    meaning = {
        "IDENTITY_MATCH_UNCERTAIN": "Vendors report different accounts for the same creator-channel and no majority agrees.",
        "IDENTITY_CONFLICT_RESOLVED": "Vendors disagreed on the account, but a majority agreed; the odd one out was dropped.",
        "CROSS_VENDOR_DISCREPANCY": "Vendors' values differ beyond the threshold with no credible majority; value held back.",
        "OUTLIER_SOURCE_EXCLUDED": "One vendor's value was an outlier against two that agree; the agreeing value was used.",
        "PRIMARY_SOURCE_MISSING": "Tier 1 couldn't supply the field (no record, stale, invalid or excluded); a lower tier did.",
        "HISTORICAL_ANOMALY": "Unusually large change vs the previous month; kept, since it can be genuine.",
    }
    counts = pd.Series([f.code for f in flags]).value_counts().to_dict() if flags else {}
    for i, (code, sev) in enumerate(p.cfg["qa_rules"]["severity"].items(), start=5):
        for c, v in enumerate([code, sev, counts.get(code, 0), meaning.get(code, ""), ACTIONS.get(code, "")], start=1):
            cell = wr.cell(i, c, v)
            cell.font = Font(name=FONT, bold=c == 1)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        wr.cell(i, 2).fill = SEV_FILL.get(sev, SEV_FILL["GREEN"])
    return wb


def write_qa_report(p, path: Path) -> Path:
    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    status_txt = {"awaiting_review": "AWAITING REVIEW — leaderboard not built yet", "new": "NEW",
                  "ready": "READY — nothing blocked", "finalized": "FINALIZED — leaderboard built"}[p.status]
    _title(s, f"QA report — {p.month}", status_txt)
    s["A2"].font = Font(name=FONT, bold=True, color="C2410C" if p.status == "awaiting_review" else "0A0A0A")
    s.column_dimensions["A"].width, s.column_dimensions["B"].width, s.column_dimensions["C"].width = 44, 14, 70
    df = p.flags.frame()
    status = p.flags.creator_status()
    n_req = len(p.req_creators)
    rows = [("Creators requested", n_req),
            ("Creators GREEN (no QA issue)", n_req - len(status)),
            ("Creators YELLOW (warnings only)", sum(1 for v in status.values() if v == "YELLOW")),
            ("Creators RED (something blocked)", sum(1 for v in status.values() if v == "RED")),
            ("Blocked items open", len(p.flags.open_red())),
            ("Blocked items resolved", sum(1 for f in p.flags.flags if f.severity == "RED" and f.status == "resolved")),
            ("Warnings", int((df["severity"] == "YELLOW").sum()) if len(df) else 0),
            ("Leaderboard gate", p.cfg["review"].get("leaderboard_gate", "strict"))]
    for i, (k, v) in enumerate(rows, start=4):
        s.cell(i, 1, k).font = Font(name=FONT, bold=True)
        s.cell(i, 2, v).font = Font(name=FONT)
    k = 4 + len(rows) + 1
    s.cell(k, 1, "By reason code").font = Font(name=FONT, bold=True)
    _head(s, k + 1, ["Reason code", "Severity", "Count"], [44, 14, 12])
    if len(df):
        for j, ((code, sev), n) in enumerate(df.groupby(["code", "severity"]).size().items(), start=k + 2):
            s.cell(j, 1, code); s.cell(j, 2, sev).fill = SEV_FILL[sev]; s.cell(j, 3, int(n))
        k = j + 2
    else:
        k += 3
    s.cell(k, 1, "How review works").font = Font(name=FONT, bold=True)
    steps = ["1. Open the 'Blocked — review' tab. Each row is one held-back creator-channel or metric.",
             "2. Read the vendor values, explanation and recommended action; pick a Decision (and a Manual value if MANUAL_VALUE).",
             "3. Add your name in Reviewer. Leave Decision empty to keep an item open.",
             f"4. Run: python review.py --month {p.month} --decisions <this file>  (or upload it on the app).",
             "5. Only the affected records are reprocessed. The leaderboard is built once no blocked item is open.",
             "Warnings (YELLOW) never block anything; they're listed for visibility."]
    for j, t in enumerate(steps, start=k + 1):
        s.cell(j, 1, t).font = Font(name=FONT)
    write_qa_sheets(wb, p, review_form=p.status != "finalized")
    if p.status == "finalized":
        write_audit_sheet(wb, p)
    wb.save(path)
    return Path(path)


def write_audit_sheet(wb, p):
    path = Path(p.registry.path).parent / "review_audit.csv"
    wa = wb.create_sheet("Review Audit")
    _title(wa, "Review audit trail (this month)", "Every human decision with the original values and what was accepted.")
    hdr = ["Reviewed at (UTC)", "Review ID", "Creator ID", "Channel", "Field(s)", "Reason code", "Original values",
           "Decision", "Manual value", "Final value", "Final source", "Reviewer", "Note"]
    _head(wa, 4, hdr, [20, 9, 11, 10, 22, 26, 50, 18, 12, 14, 18, 16, 30])
    if path.exists():
        a = pd.read_csv(path)
        a = a[a["month"].astype(str) == p.month]
        for i, r in enumerate(a.to_dict("records"), start=5):
            for c, k in enumerate(["reviewed_at", "review_id", "creator_id", "platform", "fields", "code", "original_values",
                                   "decision", "manual_value", "final_value", "final_source", "reviewer", "note"], start=1):
                v = r.get(k)
                wa.cell(i, c, None if pd.isna(v) else v).font = Font(name=FONT)


# ------------------------------------------------------------------------------------------ apply decisions
def read_decisions(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, data_only=True)
    ws = wb["Blocked — review"]
    hdr = [c.value for c in ws[4]]
    rows = [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=5)]
    df = pd.DataFrame(rows)
    if df.empty or "Decision" not in df:
        return pd.DataFrame(columns=["Review ID", "Decision"])
    return df[df["Decision"].notna() & (df["Decision"].astype(str).str.strip() != "")]


def _vendor_of(decision: str, vendors: list) -> str | None:
    for v in vendors:
        if decision == f"SELECT_{v.upper()}":
            return v
    return None


def _refill_pair(p, cid, plat, vendors_in_order):
    """Re-run the tier waterfall for one creator-channel using the given vendors."""
    for col in p.ch_values:
        p.out_ch.at[(cid, plat), col] = pd.NA
        p.out_ch.at[(cid, plat), f"{col}__source"] = pd.NA
    for v in vendors_in_order:
        ch = p.prepared[v][1]
        if (cid, plat) not in ch.index:
            continue
        for col in p.ch_values:
            if pd.isna(p.out_ch.at[(cid, plat), col]) and pd.notna(ch.at[(cid, plat), col]):
                p.out_ch.at[(cid, plat), col] = ch.at[(cid, plat), col]
                p.out_ch.at[(cid, plat), f"{col}__source"] = v
        for m in p.metrics:
            if pd.isna(p.out_ch.at[(cid, plat), f"{m}__source"]) and ch.at[(cid, plat), f"_diag_{m}"] == "likes hidden":
                p.out_ch.at[(cid, plat), f"{m}__source"] = f"n/a (likes hidden, {v})"


def apply_decisions(p, decisions: pd.DataFrame, reviewer: str = "") -> list[str]:
    """Apply reviewer decisions to the affected records only. Returns problems (rows left open)."""
    from qa_checks import differs, discrepancy_check
    vendors = list(p.cfg["vendors"])
    allowed = p.cfg["review"]["decisions"]
    by_id = {f.id: f for f in p.flags.flags}
    problems, audit = [], []
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    for r in decisions.to_dict("records"):
        f = by_id.get(str(r.get("Review ID")))
        raw = r.get("Decision")
        if raw is None or (isinstance(raw, float) and pd.isna(raw)) or str(raw).strip() == "":
            continue                                             # no decision yet: the item simply stays open
        dec = str(raw).strip().upper()
        if f is None or f.severity != "RED":
            problems.append(f"{r.get('Review ID')}: not a blocked item"); continue
        if f.status != "open":
            continue
        if dec not in allowed.get(f.code, []):
            problems.append(f"{f.id}: '{dec}' isn't allowed for {f.code} (allowed: {', '.join(allowed.get(f.code, []))})"); continue
        pair = (f.creator_id, f.platform)
        manual = r.get("Manual value")
        v_sel = _vendor_of(dec, vendors)
        cols = [f.field, *f.dependents]
        if f.code == "IDENTITY_MATCH_UNCERTAIN":
            if v_sel and v_sel not in f.values:
                problems.append(f"{f.id}: {v_sel} has no record for this channel"); continue
            if dec in ("CONFIRM_MATCH",) or v_sel:
                keep = [v for v in p.tiers if (v in f.values) and (dec == "CONFIRM_MATCH" or v == v_sel)]
                for v in f.values:
                    p.excluded_rows.discard((v, *pair))
                    if v not in keep:
                        p.excluded_rows.add((v, *pair))
                p.blocked_pairs.discard(pair)
                p.blocked_fields.pop(pair, None)
                for v in f.values:
                    p.prepared[v] = p._prepare(v)
                _refill_pair(p, *pair, keep)
                if dec == "CONFIRM_MATCH":                       # merged sources must still agree on the numbers
                    discrepancy_check(p, only_pairs={pair})
                final_src = ", ".join(keep)
                final_val = f"@{f.values[keep[0]]}" if keep else None
            else:                                                # REJECT_MATCH / EXCLUDE: channel stays out
                for col in p.ch_values:
                    p.out_ch.at[pair, f"{col}__source"] = f"EXCLUDED: {f.id}"
                final_src, final_val = "excluded", None
        else:                                                    # CROSS_VENDOR_DISCREPANCY
            if dec == "MANUAL_VALUE":
                try:
                    manual = float(manual)
                except (TypeError, ValueError):
                    problems.append(f"{f.id}: MANUAL_VALUE needs a number in 'Manual value'"); continue
            if v_sel and v_sel not in f.values:
                problems.append(f"{f.id}: {v_sel} didn't report {f.field} for this channel"); continue
            rule = p.cfg["qa_rules"]["discrepancy"]["fields"][f.field]
            if dec == "EXCLUDE":
                for col in cols:
                    p.out_ch.at[pair, col] = pd.NA
                    p.out_ch.at[pair, f"{col}__source"] = f"EXCLUDED: {f.id}"
                final_src, final_val = "excluded", None
            else:
                target = manual if dec == "MANUAL_VALUE" else f.values[v_sel]
                p.out_ch.at[pair, f.field] = int(target) if f.field == "follower_count" else target
                p.out_ch.at[pair, f"{f.field}__source"] = "manual (review)" if dec == "MANUAL_VALUE" else v_sel
                # dependent metrics: from the chosen vendor, else from a vendor whose value agrees with the target
                donors = ([v_sel] if v_sel else []) + sorted(
                    [v for v in f.values if differs(f.values[v], target, rule) <= rule["threshold"]],
                    key=lambda v: abs(f.values[v] - target))
                for col in f.dependents:
                    p.out_ch.at[pair, col] = pd.NA
                    p.out_ch.at[pair, f"{col}__source"] = f"EXCLUDED: {f.id} (no consistent source)"
                    for v in donors:
                        val = p.prepared[v][1].at[pair, col] if pair in p.prepared[v][1].index else None
                        if val is not None and pd.notna(val):
                            p.out_ch.at[pair, col], p.out_ch.at[pair, f"{col}__source"] = val, v
                            break
                final_src = p.out_ch.at[pair, f"{f.field}__source"]
                final_val = target
            fb = p.blocked_fields.get(pair, {})
            for col in cols:
                fb.pop(col, None)
        f.status, f.decision, f.manual_value = "resolved", dec, (manual if dec == "MANUAL_VALUE" else None)
        f.final_value, f.final_source = final_val, str(final_src)
        f.reviewer = str(r.get("Reviewer") or reviewer or "")
        f.note = "" if r.get("Note") is None or pd.isna(r.get("Note")) else str(r.get("Note"))
        f.reviewed_at = now
        audit.append({"month": p.month, "reviewed_at": now, "review_id": f.id, "creator_id": f.creator_id,
                      "platform": f.platform, "fields": ", ".join(cols), "code": f.code,
                      "original_values": json.dumps(f.values), "decision": dec, "manual_value": f.manual_value,
                      "final_value": final_val, "final_source": str(final_src), "reviewer": f.reviewer, "note": f.note})
    if audit:
        path = Path(p.registry.path).parent / "review_audit.csv"
        old = pd.read_csv(path) if path.exists() else pd.DataFrame()
        pd.concat([old, pd.DataFrame(audit)], ignore_index=True).to_csv(path, index=False)
    p.final = p._assess()
    p.agent.decide("review", {"applied": len(audit), "problems": problems}, f"{len(audit)} decision(s) applied",
                   f"{len(p.flags.open_red())} blocked item(s) still open" + (f"; {len(problems)} row(s) need fixing" if problems else ""))
    return problems


def resume_df(p, decisions: pd.DataFrame, reviewer: str = "", finalize: bool = True):
    """Apply decisions given as a table (columns: Review ID, Decision, Manual value, Reviewer, Note)."""
    problems = apply_decisions(p, decisions, reviewer)
    p.status = "ready" if p.gate_open() else "awaiting_review"
    p.flags.frame().to_csv(p.out_dir / "qa_flags.csv", index=False)
    if p.status == "ready" and finalize:
        p.finalize()
    else:
        write_qa_report(p, p.out_dir / f"qa_report_{p.month}.xlsx")
        p.save_state()
    return p, problems


def resume(month: str, decisions_path: Path, reviewer: str = "", out_dir: Path | None = None):
    p = __import__("pipeline").Pipeline.load_state(month, out_dir)
    problems = apply_decisions(p, read_decisions(decisions_path), reviewer)
    p.status = "ready" if p.gate_open() else "awaiting_review"
    p.flags.frame().to_csv(p.out_dir / "qa_flags.csv", index=False)
    if p.status == "ready":
        p.finalize()
    else:
        write_qa_report(p, p.out_dir / f"qa_report_{p.month}.xlsx")
        p.save_state()
        print(f"{len(p.flags.open_red())} blocked item(s) still open; updated {p.out_dir / f'qa_report_{p.month}.xlsx'}")
    for pr in problems:
        print("  needs fixing:", pr)
    return p, problems


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", required=True)
    ap.add_argument("--decisions", required=True, help="the QA report with the Decision column filled in")
    ap.add_argument("--reviewer", default="")
    a = ap.parse_args()
    resume(a.month, Path(a.decisions), a.reviewer)
