"""Write the run results as one formatted Excel workbook.

Sheets: Summary · Leaderboard · Instagram · TikTok · YouTube · Channels · Creators · Metric Definitions · Vendor QA · Vendor History · Creator History · Waterfall · Gaps · Decision Log · Blocked — review · Warnings · Reason codes · Review Audit
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="0A0A0A")
PILLAR_FILL = PatternFill("solid", start_color="EA580C")          # orange header for the three pillars
TIER_FILL = {1: PatternFill("solid", start_color="FDBA74"),
             2: PatternFill("solid", start_color="FED7AA"),
             3: PatternFill("solid", start_color="FFF1E6")}
EMPTY_FILL = PatternFill("solid", start_color="F2F2F2")
INPUT_FONT = Font(name=FONT, color="0000FF")
MUTED = Font(name=FONT, color="8C8C8C", italic=True)
BOLD = Font(name=FONT, bold=True)
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
NUMFMT = {"multiple": '0.00"x"', "percent_signed": '+0.0%;-0.0%;0.0%', "percent": "0.0%"}


def _v(x):
    return None if x is None or x is pd.NA or (isinstance(x, float) and pd.isna(x)) else x


def _header(ws, row, headers, widths=None, fills=None):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = (fills or {}).get(c, HEAD_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    for c, w in enumerate(widths or [], start=1):
        ws.column_dimensions[L(c)].width = w


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color="0A0A0A")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, italic=True, color="595959")


def _fonts(wb):
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                f = cell.font
                if f is None or f.name != FONT:
                    cell.font = Font(name=FONT, bold=f.bold if f else False, italic=f.italic if f else False,
                                     color=f.color if f else None)


def write_workbook(p, path: Path):
    tiers = {name: i for i, name in enumerate(p.tiers, start=1)}
    fill_for = lambda src: TIER_FILL.get(tiers.get(src), EMPTY_FILL) if _v(src) else EMPTY_FILL
    metrics, cr_fields, ch_fields = p.metrics, list(p.cr_fields), list(p.ch_fields)
    mkeys = list(metrics)
    wb = Workbook()

    # ------------------------------------------------------------------ Leaderboard
    lb = p.leaderboard
    ws = wb.active
    ws.title = "Leaderboard"
    _title(ws, "Creator Leaderboard — The Pulse · The Lens · The Hub",
           "Scores are 0–100 percentiles: each channel is compared with requested channels on the same platform "
           "(growth also within follower-size band), then averaged over the creator's requested channels. "
           "Blank Hub score = likes hidden on every channel.")
    hdr = ["Creator", "Creator ID", "Country", "Category", "Channels", "Total followers"]
    fills = {}
    for m in mkeys:
        hdr += [f"{metrics[m]['pillar']} score", f"{metrics[m]['pillar']} rank"]
        fills[len(hdr) - 1] = fills[len(hdr)] = PILLAR_FILL
    hdr += ["Primary role", "Confidence", "Sources"]
    _header(ws, 4, hdr, [18, 11, 9, 12, 24, 14] + [11, 10] * len(mkeys) + [13, 11, 22], fills)
    first, last = 5, 4 + len(lb)
    sc = {m: 7 + 2 * j for j, m in enumerate(mkeys)}          # score column per metric
    for i, r in enumerate(lb.to_dict("records"), start=first):
        for c, v in enumerate([r["display_name"], r["creator_id"], r["country"], _v(r["category"]),
                               ", ".join(r["channels"]), r["total_followers"]], start=1):
            ws.cell(i, c, v)
        ws.cell(i, 6).number_format = "#,##0"
        for m in mkeys:
            col = L(sc[m])
            ws.cell(i, sc[m], _v(r[m])).number_format = "0"
            ws.cell(i, sc[m] + 1, f'=IF({col}{i}="","",RANK({col}{i},${col}${first}:${col}${last}))')
        a, b, c3 = (f"N({L(sc[m])}{i})" for m in mkeys)
        names = [metrics[m]["pillar"] for m in mkeys]
        rc = 7 + 2 * len(mkeys)
        ws.cell(i, rc, f'=IF(AND({a}>={b},{a}>={c3}),"{names[0]}",IF({b}>={c3},"{names[1]}","{names[2]}"))')
        conf = ws.cell(i, rc + 1, r["confidence"])
        conf.fill = {"High": TIER_FILL[1], "Medium": TIER_FILL[2]}.get(r["confidence"], TIER_FILL[3])
        ws.cell(i, rc + 2, r["sources"])
        for c in range(1, rc + 3):
            ws.cell(i, c).border = THIN
    if len(lb):
        for m in mkeys:
            rng = f"{L(sc[m])}{first}:{L(sc[m])}{last}"
            ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                              end_type="num", end_value=100, end_color="F97316"))
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{L(len(hdr))}{max(last, 5)}"

    # ------------------------------------------------------------------ Per-platform leaderboards
    cb = p.channel_board
    for pl in p.platforms:
        rows = cb[cb["platform"] == pl].sort_values(f"{mkeys[0]}__pct", ascending=False)
        if rows.empty:
            continue
        wp = wb.create_sheet(pl.capitalize().replace("tok", "Tok").replace("tube", "Tube"))
        _title(wp, f"{wp.title} leaderboard — one row per complete {wp.title} channel",
               "Ranks every requested channel on this platform that is filled, even if the creator's other channels "
               "aren't yet. Scores are 0–100 percentiles within the platform (The Lens also within size band).")
        hdr = ["Creator", "Creator ID", "Handle", "Country", "Followers", "Size band"]
        fills = {}
        col_of = {}
        for m in mkeys:
            hdr += [metrics[m]["label"], f"{metrics[m]['pillar']} score", f"{metrics[m]['pillar']} rank"]
            col_of[m] = len(hdr) - 1
            for c in range(len(hdr) - 2, len(hdr) + 1):
                fills[c] = PILLAR_FILL
        hdr += ["Primary role", "Confidence", "All requested channels filled", "Sources"]
        _header(wp, 4, hdr, [18, 11, 24, 9, 12, 10] + [13, 10, 9] * len(mkeys) + [13, 11, 13, 22], fills)
        first, last = 5, 4 + len(rows)
        for i, r in enumerate(rows.to_dict("records"), start=first):
            for c, v in enumerate([_v(r.get("display_name")) or r[p.key], r[p.key], r["handle"], _v(r.get("country")),
                                   int(r["follower_count"]), r["size_band"]], start=1):
                wp.cell(i, c, v)
            wp.cell(i, 5).number_format = "#,##0"
            for m in mkeys:
                c = col_of[m]
                if str(r[f"{m}__source"]).startswith("n/a"):
                    wp.cell(i, c - 1, "n/a – likes hidden").font = MUTED
                else:
                    wp.cell(i, c - 1, _v(r[m])).number_format = NUMFMT[metrics[m]["display"]]
                wp.cell(i, c, _v(r[f"{m}__pct"])).number_format = "0"
                cl = L(c)
                wp.cell(i, c + 1, f'=IF({cl}{i}="","",RANK({cl}{i},${cl}${first}:${cl}${last}))')
            sa, sb, sc3 = (f"N({L(col_of[m])}{i})" for m in mkeys)
            nm = [metrics[m]["pillar"] for m in mkeys]
            rcol = 7 + 3 * len(mkeys)
            wp.cell(i, rcol, f'=IF(AND({sa}>={sb},{sa}>={sc3}),"{nm[0]}",IF({sb}>={sc3},"{nm[1]}","{nm[2]}"))')
            cf = wp.cell(i, rcol + 1, r["confidence"])
            cf.fill = {"High": TIER_FILL[1], "Medium": TIER_FILL[2]}.get(r["confidence"], TIER_FILL[3])
            done = bool(r["creator_complete"])
            wp.cell(i, rcol + 2, "Yes" if done else "No").font = Font(name=FONT, bold=True,
                                                                     color="0A0A0A" if done else "C2410C")
            wp.cell(i, rcol + 3, r["sources"])
            for c in range(1, rcol + 4):
                wp.cell(i, c).border = THIN
        for m in mkeys:
            rng = f"{L(col_of[m])}{first}:{L(col_of[m])}{last}"
            wp.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                              end_type="num", end_value=100, end_color="F97316"))
        wp.freeze_panes = "C5"
        wp.auto_filter.ref = f"A4:{L(len(hdr))}{last}"

    # ------------------------------------------------------------------ Channels
    wh = wb.create_sheet("Channels")
    _title(wh, "Channel metrics (one row per requested creator + channel)",
           "Value colour = tier that supplied it. Each metric is computed from a single vendor record. "
           "Percentile = rank among requested channels in the peer group shown.")
    ch = p.scored_ch
    hdr = ["Creator ID", "Channel", "Handle", "Followers", "Size band"]
    fills = {}
    for m in mkeys:
        hdr += [metrics[m]["label"], f"{metrics[m]['pillar']} percentile"]
        fills[len(hdr) - 1] = fills[len(hdr)] = PILLAR_FILL
        if metrics[m].get("compare_within_size_band"):
            hdr.append(f"{metrics[m]['pillar']} peer group")
            fills[len(hdr)] = PILLAR_FILL
    hdr += ["Complete"] + [f"{c} source" for c in ["follower_count", *mkeys]]
    widths = []
    for h in hdr:
        widths.append(22 if h == "Handle" else 18 if h.endswith("peer group") else 16 if h.endswith(" source")
                      else 13 if h.startswith("90d") or h.endswith("percentile") else 10 if h in ("Complete", "Size band")
                      else 11)
    _header(wh, 4, hdr, widths, fills)
    filled = pd.concat([p.filled(ch, f) for f in p.ch_required], axis=1).all(axis=1)
    for i, ((k, pl), r) in enumerate(ch.iterrows(), start=5):
        wh.cell(i, 1, k); wh.cell(i, 2, pl)
        wh.cell(i, 3, _v(r["handle"])).fill = fill_for(r["handle__source"])
        fc = wh.cell(i, 4, None if _v(r["follower_count"]) is None else int(float(r["follower_count"])))
        fc.fill, fc.number_format = fill_for(r["follower_count__source"]), "#,##0"
        wh.cell(i, 5, _v(r["size_band"]))
        c = 6
        for m in mkeys:
            src = r[f"{m}__source"]
            if str(src).startswith("n/a"):
                wh.cell(i, c, "n/a – likes hidden").font = MUTED
            else:
                cell = wh.cell(i, c, _v(r[m]))
                cell.fill, cell.number_format = fill_for(src), NUMFMT[metrics[m]["display"]]
            wh.cell(i, c + 1, _v(r[f"{m}__pct"])).number_format = "0"
            c += 2
            if metrics[m].get("compare_within_size_band"):
                wh.cell(i, c, _v(r[f"{m}__peers"]) or None)
                c += 1
        ok = bool(filled.at[(k, pl)])
        wh.cell(i, c, "Yes" if ok else "No").font = Font(name=FONT, bold=True, color="0A0A0A" if ok else "C2410C")
        for j, f in enumerate(["follower_count", *mkeys], start=c + 1):
            s = _v(r[f"{f}__source"])
            wh.cell(i, j, s)
    wh.freeze_panes = "C5"
    wh.auto_filter.ref = f"A4:{L(len(hdr))}{4 + len(ch)}"

    # ------------------------------------------------------------------ Creators
    wc = wb.create_sheet("Creators")
    _title(wc, "Creator details (one row per requested creator)",
           "Cell colour = tier that supplied the value: deep orange T1 · light orange T2 · pale orange T3 · grey still missing")
    per = pd.DataFrame({"n": filled.groupby(level=0).size(), "ok": filled.groupby(level=0).sum()})
    chans = ch.reset_index().groupby(p.key)["platform"].apply(lambda s: ", ".join(x for x in p.platforms if x in set(s)))
    hdr = ["Creator ID"] + [f + (" *" if f in p.required else "") for f in cr_fields] + \
          ["Channels requested", "Channels complete", "Creator complete"] + [f"{f} source" for f in cr_fields]
    _header(wc, 4, hdr, [11] + [24 if f == "contact_email" else 15 for f in cr_fields] + [24, 11, 11] + [12] * len(cr_fields))
    for i, (k, r) in enumerate(p.out_cr.iterrows(), start=5):
        wc.cell(i, 1, k)
        for c, f in enumerate(cr_fields, start=2):
            wc.cell(i, c, _v(r[f])).fill = fill_for(r[f"{f}__source"])
        base = 2 + len(cr_fields)
        wc.cell(i, base, chans.get(k))
        wc.cell(i, base + 1, f"{int(per.at[k, 'ok'])}/{int(per.at[k, 'n'])}")
        done = all(_v(r[f]) is not None for f in p.cr_required) and per.at[k, "ok"] == per.at[k, "n"]
        wc.cell(i, base + 2, "Yes" if done else "No").font = Font(name=FONT, bold=True,
                                                                  color="0A0A0A" if done else "C2410C")
        for c, f in enumerate(cr_fields, start=base + 3):
            wc.cell(i, c, _v(r[f"{f}__source"]))
    wc.freeze_panes = "B5"
    wc.auto_filter.ref = f"A4:{L(len(hdr))}{4 + len(p.out_cr)}"

    # ------------------------------------------------------------------ Metric Definitions
    wm = wb.create_sheet("Metric Definitions")
    _title(wm, "The three metrics and how their caveats are handled")
    _header(wm, 4, ["Pillar", "Theme", "Who it finds", "Metric", "Formula", "Compared within", "Plausible range",
                    "Caveat and how it is handled"], [12, 11, 30, 22, 44, 22, 14, 60])
    caveat = {
        "reach_rate": "Favours accounts that publish video. Each channel is compared only with others on the same "
                      "platform (TikTok views run far higher than Instagram), so an IG account is judged against IG.",
        "growth_rate": "Favours small accounts gaining traction. Compared within platform AND follower-size band "
                       f"({', '.join(f'{b} {v:,}+' for b, v in p.cfg['size_bands'].items())}); bands with fewer than "
                       f"{p.cfg.get('min_peers_for_band', 8)} peers fall back to the whole platform.",
        "meaningful_eng_rate": "Hidden likes would inflate the rate, so it is not computed for those accounts: "
                               "marked n/a, not counted as missing, and the creator's Hub score uses their other "
                               "channels. YouTube doesn't publish saves, so its formula leaves saves out.",
    }
    for i, (m, s) in enumerate(metrics.items(), start=5):
        f = s["formula"] + "".join(f"\n{pl}: {fx}" for pl, fx in (s.get("by_platform") or {}).items())
        within = "platform + size band" if s.get("compare_within_size_band") else "platform"
        lo, hi = s.get("bounds", ["", ""])
        for c, v in enumerate([s["pillar"], s["theme"], s["blurb"], s["label"], f, within, f"{lo} to {hi}",
                               caveat.get(m, "")], start=1):
            cell = wm.cell(i, c, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        wm.row_dimensions[i].height = 75
    n = 5 + len(metrics) + 1
    notes = ["All inputs are 90-day totals from the vendor record. A metric's inputs always come from the same vendor "
             "record; the metric itself is then filled tier by tier.",
             "Values outside the plausible range are treated as bad vendor data and rejected (counted in Vendor QA → validity).",
             "Percentiles are relative to this run's requested channels, so scores compare creators within a request, "
             "not against an external benchmark."]
    for j, t in enumerate(notes):
        c = wm.cell(n + j, 1, "• " + t)
        c.alignment = Alignment(wrap_text=False)

    # ------------------------------------------------------------------ Vendor QA
    wq = wb.create_sheet("Vendor QA")
    _title(wq, "Vendor QA scorecard — measured on requested channels only",
           "This month's score = SUMPRODUCT(weights, metrics). Blended = (1 − weight) × this month + weight × "
           "3-month vendor reliability. Tier uses the blended score. Blue cells are inputs from config.yaml; "
           "changing them re-ranks this sheet only. Rerun the pipeline to refill the data.")
    qm = list(p.cfg["qa"]["weights"])
    wq["A4"], wq["A5"] = "Weights →", "Quality floor →"
    for j, m in enumerate(qm):
        c = wq.cell(4, 2 + j, p.cfg["qa"]["weights"][m])
        c.font, c.number_format = INPUT_FONT, "0%"
    wq.cell(4, 2 + len(qm), f"=SUM(B4:{L(1 + len(qm))}4)").number_format = "0%"
    wq.cell(4, 3 + len(qm), "← must be 100%")
    wq["B5"].value, wq["B5"].font, wq["B5"].number_format = p.cfg["qa"]["min_score_to_use"], INPUT_FONT, "0.00"
    wq["A6"] = "History weight →"
    wq["B6"].value, wq["B6"].font, wq["B6"].number_format = p.cfg["qa"].get("history_weight", 0.15), INPUT_FONT, "0%"
    wq["C6"] = (f"3-month reliability = {p.cfg['qa']['reliability']['qa_score_weight']:.0%} × average QA score "
                f"+ {p.cfg['qa']['reliability']['review_accuracy_weight']:.0%} × review accuracy over the last "
                f"{p.cfg['qa']['reliability']['window_months']} months (QA average alone if no disputes were reviewed)")
    wq["B4"].comment = Comment("Source: config.yaml → qa.weights", "pipeline")
    extra = ["QA average (3 mo)", "Review accuracy (3 mo)", "Disputes reviewed (3 mo)",
             "Vendor rows", "Rows not needed", "Duplicate rate", "Implausible values", "Format fixes"] + \
            [f"{metrics[m]['pillar']} computable" for m in mkeys]
    heads = ["Vendor"] + [m.replace("_", " ").capitalize() for m in qm] + \
            ["This month's score", "3-month reliability", "Months in window", "Blended score", "Tier"] + extra
    _header(wq, 7, heads, [14] + [13] * len(qm) + [11, 11, 10, 10, 9] + [11] * len(extra))
    names = sorted(p.profiles, key=lambda n: -p.profiles[n].score)
    r0, r1 = 8, 7 + len(names)
    lm = L(1 + len(qm))
    cur, hist, nh, bl = (L(2 + len(qm) + k) for k in range(4))
    for i, n in enumerate(names, start=r0):
        pr = p.profiles[n]
        wq.cell(i, 1, n)
        for j, m in enumerate(qm):
            wq.cell(i, 2 + j, getattr(pr, m)).number_format = "0.0%"
        wq.cell(i, 2 + len(qm), f"=SUMPRODUCT($B$4:${lm}$4,B{i}:{lm}{i})").number_format = "0.000"
        wq.cell(i, 3 + len(qm), pr.history_score).number_format = "0.000"
        wq.cell(i, 4 + len(qm), pr.history_months)
        wq.cell(i, 5 + len(qm), f'=IF({hist}{i}="",{cur}{i},(1-$B$6)*{cur}{i}+$B$6*{hist}{i})').number_format = "0.000"
        wq.cell(i, 6 + len(qm), f'=IF({bl}{i}<$B$5,"excluded",COUNTIFS(${bl}${r0}:${bl}${r1},">"&{bl}{i},'
                                f'${bl}${r0}:${bl}${r1},">="&$B$5)+1)')
        vals = [(pr.qa_avg_3m, "0.000"), (pr.review_accuracy, "0%"), (pr.disputes_reviewed, "0"),
                (pr.rows, "#,##0"), (pr.rows_not_needed, "#,##0"), (pr.duplicate_rate, "0.0%"),
                (pr.implausible_values, "#,##0"), (pr.format_fixes, "#,##0")] + \
               [(pr.field_completeness[m], "0%") for m in mkeys]
        for k, (v, fmt) in enumerate(vals):
            wq.cell(i, 7 + len(qm) + k, v).number_format = fmt
    rr = r1 + 3
    wq.cell(rr, 1, "Fresh coverage by channel (share of requested channels of that type)").font = BOLD
    _header(wq, rr + 1, ["Channel", "Requested"] + names)
    plats = [x for x in p.platforms if p.request_mix.get(x)]
    for k, pl in enumerate(plats, start=rr + 2):
        wq.cell(k, 1, pl); wq.cell(k, 2, p.request_mix[pl])
        best = max(p.profiles[x].platform_coverage.get(pl, 0) for x in names)
        for j, n in enumerate(names):
            v = p.profiles[n].platform_coverage.get(pl, 0)
            c = wq.cell(k, 3 + j, v)
            c.number_format = "0%"
            if v == best and v > 0:
                c.font = Font(name=FONT, bold=True, color="0A0A0A")

    # ------------------------------------------------------------------ Vendor History
    wv = wb.create_sheet("Vendor History")
    _title(wv, "Vendor quality by month (the memory behind the history weight)",
           "This month's score per vendor in each month processed so far, and the tier it got. "
           "Past months feed the blended score at the history weight on the Vendor QA tab.")
    vh = p.vendor_history.df.copy()
    months = sorted(vh["month"].unique())
    vnames = sorted(vh["vendor"].unique())
    _header(wv, 4, ["Month"] + [f"{v} score" for v in vnames] + [f"{v} tier" for v in vnames])
    for i, mth in enumerate(months, start=5):
        wv.cell(i, 1, mth).font = BOLD if mth == p.month else Font(name=FONT)
        for j, v in enumerate(vnames):
            r = vh[(vh["month"] == mth) & (vh["vendor"] == v)]
            if len(r):
                wv.cell(i, 2 + j, float(r["current_score"].iloc[0])).number_format = "0.000"
                wv.cell(i, 2 + len(vnames) + j, r["tier"].iloc[0])
    k = 6 + len(months)
    wv.cell(k, 1, "What the pipeline has learned about each vendor's files").font = BOLD
    _header(wv, k + 1, ["Vendor", "Months seen", "Platform spellings", "File-name words", "Column names learned"])
    for c, w in zip("ABCDEFG", [12, 20, 26, 22, 46, 13, 13]):
        wv.column_dimensions[c].width = w
    for i, (v, info) in enumerate(p.registry.data.items(), start=k + 2):
        learned = "; ".join(f"{l['month']}: {l['field']} ← '{l['column']}'" for l in info.get("learned", [])) or "—"
        for c, val in enumerate([v, ", ".join(info["months_seen"]), ", ".join(info["platform_values"]),
                                 ", ".join(info["filename_tokens"]), learned], start=1):
            wv.cell(i, c, val).alignment = Alignment(wrap_text=True, vertical="top")

    # ------------------------------------------------------------------ Creator History
    wch = wb.create_sheet("Creator History")
    _title(wch, "Creator scores by month (requested creators this month)",
           "Channel = that channel's 0–100 score; 'all' = the creator's average over requested channels. "
           "Blank = not requested or not filled that month.")
    hist_df = p.creator_history.df
    hist_df = hist_df[hist_df["creator_id"].isin(set(p.req_creators))].sort_values(["creator_id", "platform", "month"])
    hdr = ["Creator ID", "Channel", "Month", "Followers"] + [f"{metrics[m]['pillar']} score" for m in mkeys] + ["Complete"]
    _header(wch, 4, hdr, [11, 10, 9, 12] + [12] * len(mkeys) + [10], {5 + j: PILLAR_FILL for j in range(len(mkeys))})
    for i, r in enumerate(hist_df.to_dict("records"), start=5):
        wch.cell(i, 1, r["creator_id"]); wch.cell(i, 2, r["platform"]); wch.cell(i, 3, r["month"])
        fcount = _v(r.get("follower_count"))
        if fcount is not None:
            wch.cell(i, 4, int(float(fcount))).number_format = "#,##0"
        for j, m in enumerate(mkeys):
            v = _v(r.get(f"{m}__pct"))
            if v is not None:
                wch.cell(i, 5 + j, float(v)).number_format = "0"
        wch.cell(i, 5 + len(mkeys), "Yes" if str(r.get("complete")) == "True" else "No")
    wch.freeze_panes = "A5"
    wch.auto_filter.ref = f"A4:{L(len(hdr))}{4 + len(hist_df)}"

    # ------------------------------------------------------------------ Waterfall
    ww = wb.create_sheet("Waterfall")
    _title(ww, "Waterfall progression", "Fill rates on the requested scope after each tier (n/a counts as filled)")
    tr = p.tier_results
    _header(ww, 4, ["Measure"] + [f"After T{t['tier']} ({t['vendor']})" for t in tr]
            + [f"Cells filled by T{t['tier']}" for t in tr], [34] + [16] * (2 * len(tr)))
    k = 5
    for f in cr_fields + p.ch_values:
        lab = f"{metrics[f]['pillar']} ({metrics[f]['label']})" if f in metrics else f
        ww.cell(k, 1, lab + (" *" if f in p.required else ""))
        for j, t in enumerate(tr):
            ww.cell(k, 2 + j, t["after"]["field_fill_rate"][f]).number_format = "0.0%"
            ww.cell(k, 2 + len(tr) + j, t["cells_filled"][f]).number_format = "#,##0"
        k += 1
    lastf = k - 1
    ww.cell(k, 1, "Total cells filled").font = BOLD
    for j in range(len(tr)):
        col = L(2 + len(tr) + j)
        ww.cell(k, 2 + len(tr) + j, f"=SUM({col}5:{col}{lastf})").number_format = "#,##0"
    k += 2
    ww.cell(k, 1, "Requested channels complete, by channel").font = BOLD
    k += 1
    for pl in plats:
        ww.cell(k, 1, f"{pl} ({p.request_mix[pl]} requested)")
        for j, t in enumerate(tr):
            ww.cell(k, 2 + j, t["after"]["channels_complete_by_platform"].get(pl, 0)).number_format = "0.0%"
        k += 1
    for label, key in [("All requested channels complete", "channels_complete"),
                       ("Creators complete (details + every requested channel)", "creators_complete")]:
        ww.cell(k, 1, label).font = BOLD
        for j, t in enumerate(tr):
            ww.cell(k, 2 + j, t["after"][key]).number_format = "0.0%"
        k += 1
    ww.cell(k + 1, 1, f"Target: every required field and metric ≥ {p.cfg['waterfall']['required_field_target']:.0%}")
    ww.cell(k + 2, 1, "* required")

    # ------------------------------------------------------------------ Gaps
    wg = wb.create_sheet("Gaps")
    _title(wg, "Still missing (required items only)",
           "Reasons name the vendor: stale data, a missing input for a metric, or an implausible value that was rejected.")
    _header(wg, 4, ["Creator ID", "Level", "Channel", "Missing", "Reason"], [11, 9, 11, 40, 80])
    for i, r in enumerate(p.gaps.itertuples(index=False), start=5):
        for c, v in enumerate(r, start=1):
            wg.cell(i, c, v or None)
    wg.freeze_panes = "A5"
    wg.auto_filter.ref = f"A4:E{max(4 + len(p.gaps), 5)}"

    # ------------------------------------------------------------------ Decision Log
    wd = wb.create_sheet("Decision Log")
    _title(wd, "Agent decision log", "Every accept / flag / go / skip / stop decision with the reason behind it")
    _header(wd, 4, ["#", "Step", "Decision", "Reason"], [5, 20, 14, 110])
    for i, e in enumerate(p.agent.log, start=5):
        for c, v in enumerate([i - 4, e["step"], e["decision"], e["reason"]], start=1):
            wd.cell(i, c, v)
        wd.cell(i, 4).alignment = Alignment(wrap_text=True, vertical="top")

    # ------------------------------------------------------------------ Summary (first)
    s = wb.create_sheet("Summary", 0)
    _title(s, f"Creator vendor fill — {p.month}", f"Monthly run · freshness measured at {p.run_date.date()}")
    s.column_dimensions["A"].width, s.column_dimensions["B"].width, s.column_dimensions["C"].width = 46, 22, 50
    n_cr, n_ch = len(p.out_cr), len(p.scored_ch)
    ccol = L(4 + len(cr_fields))
    comp_col = L(6 + 2 * len(mkeys) + sum(1 for m in mkeys if metrics[m].get("compare_within_size_band")))
    rc = L(7 + 2 * len(mkeys))
    rows = [
        ("Creators requested", f"=COUNTA(Creators!A5:A{4 + n_cr})", "#,##0"),
        ("Channels requested", f"=COUNTA(Channels!A5:A{4 + n_ch})", "#,##0"),
        ("Creators complete (on leaderboard)", f'=COUNTIF(Creators!{ccol}5:{ccol}{4 + n_cr},"Yes")', "#,##0"),
        ("Share of creators complete", "=IFERROR(B6/B4,0)", "0.0%"),
        ("Requested channels complete", f'=COUNTIF(Channels!{comp_col}5:{comp_col}{4 + n_ch},"Yes")', "#,##0"),
        ("Share of channels complete", "=IFERROR(B8/B5,0)", "0.0%"),
        ("Missing items (see Gaps)", f"=COUNTA(Gaps!A5:A{max(4 + len(p.gaps), 5)})", "#,##0"),
        ("Tier order used", " → ".join(f"T{t['tier']} {t['vendor']}" for t in tr), None),
    ]
    for i, (label, val, fmt) in enumerate(rows, start=4):
        s.cell(i, 1, label).font = BOLD
        c = s.cell(i, 2, val)
        if fmt:
            c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
    k = 14
    s.cell(k, 1, "Files this month (recognised automatically)").font = BOLD
    _header(s, k + 1, ["File", "Recognised as", "Match", "Notes"])
    for j, r in enumerate(p.recognition, start=k + 2):
        notes = []
        for l in r.get("learned", []):
            notes.append(f"learned: {l['field']} is now '{l['column']}'")
        if r.get("unused"):
            notes.append("new unused column(s): " + ", ".join(r["unused"]))
        if r.get("missing"):
            notes.append("no column for: " + ", ".join(r["missing"]))
        if r["kind"] in ("unrecognised", "duplicate"):
            notes.append(r["kind"] + " — skipped")
        s.cell(j, 1, r["file"]); s.cell(j, 2, r["vendor"] or r["kind"])
        s.cell(j, 3, r["confidence"]).number_format = "0%"
        s.cell(j, 4, "; ".join(notes) or None)
    s.column_dimensions["D"].width = 60
    k = j + 2
    s.cell(k, 1, "Primary role of complete creators").font = BOLD
    _header(s, k + 1, ["Pillar", "Creators", "Who it finds"], fills={1: PILLAR_FILL, 2: PILLAR_FILL, 3: PILLAR_FILL})
    for j, m in enumerate(mkeys, start=k + 2):
        s.cell(j, 1, metrics[m]["pillar"])
        s.cell(j, 2, f'=COUNTIF(Leaderboard!{rc}5:{rc}{max(last, 5)},A{j})')
        s.cell(j, 3, f"{metrics[m]['theme']}: {metrics[m]['blurb']}")
    k = j + 2
    s.cell(k, 1, "Creators by number of channels requested").font = BOLD
    _header(s, k + 1, ["Channels requested", "Creators", "Complete", "Share complete"])
    chc, donec = L(3 + len(cr_fields)), L(4 + len(cr_fields))
    rng_ch, rng_done = f"Creators!${chc}$5:${chc}${4 + n_cr}", f"Creators!${donec}$5:${donec}${4 + n_cr}"
    for j, n in enumerate(sorted(p.request_shape, reverse=True), start=k + 2):
        s.cell(j, 1, f"{n} of {len(p.platforms)}")
        s.cell(j, 2, f'=COUNTIF({rng_ch},"*/{n}")').number_format = "#,##0"
        s.cell(j, 3, f'=COUNTIFS({rng_ch},"*/{n}",{rng_done},"Yes")').number_format = "#,##0"
        s.cell(j, 4, f"=IFERROR(C{j}/B{j},0)").number_format = "0.0%"
    k = j + 2
    s.cell(k, 1, "Requested channels").font = BOLD
    _header(s, k + 1, ["Channel", "Requested", "Complete after last tier"])
    for j, pl in enumerate(plats, start=k + 2):
        s.cell(j, 1, pl); s.cell(j, 2, p.request_mix[pl])
        s.cell(j, 3, p.final["channels_complete_by_platform"].get(pl, 0)).number_format = "0.0%"
    k = j + 2
    s.cell(k, 1, "Colour key").font = BOLD
    for j, (lbl, fl) in enumerate([("Tier 1 value", TIER_FILL[1]), ("Tier 2 value", TIER_FILL[2]),
                                   ("Tier 3 value", TIER_FILL[3]), ("Still missing", EMPTY_FILL)], start=k + 1):
        s.cell(j, 1, lbl).fill = fl

    from review import write_audit_sheet, write_qa_sheets
    write_qa_sheets(wb, p, review_form=False)
    write_audit_sheet(wb, p)
    _fonts(wb)
    wb.save(path)
    return path
