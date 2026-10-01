"""Creator request list: the Excel template people fill in, and writing a filled request file."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
HEAD = PatternFill("solid", start_color="0A0A0A")
INPUT = PatternFill("solid", start_color="FFFDE7")
EXAMPLE = PatternFill("solid", start_color="FFF1E6")


def _head(ws, row, values, widths):
    for c, (v, w) in enumerate(zip(values, widths), start=1):
        cell = ws.cell(row, c, v)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center" if c > 2 else "left")
        ws.column_dimensions[cell.column_letter].width = w


def write_request(path: Path, platforms: list[str], rows: list[dict] | None = None, blank_rows: int = 500) -> Path:
    """rows: [{"creator_id": .., "creator_name": .., "channels": ["instagram", ...]}]. None -> empty template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Request"
    headers = ["creator_id", "creator_name (optional)"] + platforms
    _head(ws, 1, headers, [14, 26] + [12] * len(platforms))
    ws.freeze_panes = "C2"
    last_col = ws.cell(1, len(headers)).column_letter
    n = max(len(rows or []), 0) + blank_rows
    dv = DataValidation(type="list", formula1='"Y"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Use Y or leave blank", error="Mark a needed channel with Y; leave it blank otherwise.")
    ws.add_data_validation(dv)
    dv.add(f"C2:{last_col}{n + 1}")
    for i, r in enumerate(rows or [], start=2):
        ws.cell(i, 1, r["creator_id"])
        ws.cell(i, 2, r.get("creator_name"))
        for j, p in enumerate(platforms, start=3):
            if p in r["channels"]:
                ws.cell(i, j, "Y")
    for row in ws.iter_rows(min_row=2, max_row=n + 1, max_col=len(headers)):
        for cell in row:
            cell.font = Font(name=FONT, color="0000FF")
            if cell.column > 2:
                cell.alignment = Alignment(horizontal="center")

    # ---- How to fill
    h = wb.create_sheet("How to fill")
    h.column_dimensions["A"].width = 100
    lines = [
        ("Creator request list — how to fill", Font(name=FONT, bold=True, size=14, color="0A0A0A")),
        ("Upload this file with every run. It tells the pipeline which creators, and which of their channels, we need.", None),
        ("", None),
        ("1.  One row per creator on the Request sheet. creator_id must match the ID the vendors use.", None),
        ("2.  Put Y under each channel you need for that creator. Leave the other channels blank.", None),
        ("3.  creator_name is only for your reference; the pipeline fills the name from vendor data.", None),
        ("", None),
        ("Why it matters: completeness and vendor quality are measured only on the channels you mark.", Font(name=FONT, bold=True)),
        ("If creator A is marked for Instagram and TikTok, a vendor with no YouTube data for A is not penalised, "
         "and A counts as complete once Instagram and TikTok are filled.", None),
        ("", None),
        ("Example (not read by the pipeline):", Font(name=FONT, bold=True)),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        c = h.cell(i, 1, t)
        c.font = f or Font(name=FONT)
        c.alignment = Alignment(wrap_text=True)
    start = len(lines) + 1
    ex_headers = headers
    for c, v in enumerate(ex_headers, start=1):
        cell = h.cell(start, c, v)
        cell.font, cell.fill = Font(name=FONT, bold=True, color="FFFFFF"), HEAD
    examples = [("CR00012", "Creator A", {"instagram", "tiktok"}, "needs IG + TikTok only; YouTube gaps are ignored"),
                ("CR00345", "Creator B", {"youtube"}, "YouTube only"),
                ("CR01020", "Creator C", {"youtube", "instagram", "tiktok"}, "all three channels")]
    for k, (cid, name, chans, note) in enumerate(examples, start=start + 1):
        for c, v in enumerate([cid, name] + ["Y" if p in chans else None for p in platforms], start=1):
            cell = h.cell(k, c, v)
            cell.font, cell.fill = Font(name=FONT), EXAMPLE
        h.cell(k, len(ex_headers) + 1, note).font = Font(name=FONT, italic=True, color="595959")
    for col in range(2, len(ex_headers) + 1):
        h.column_dimensions[L(col)].width = 12 if col > 2 else 24
    h.column_dimensions["A"].width = 14
    h.column_dimensions[L(len(ex_headers) + 1)].width = 44
    h.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(ex_headers) + 1)
    for r in range(2, len(lines) + 1):
        h.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(ex_headers) + 1)
        h.row_dimensions[r].height = 30 if len(str(h.cell(r, 1).value or "")) > 95 else 15
    wb.save(path)
    return Path(path)


if __name__ == "__main__":
    import yaml
    root = Path(__file__).parent
    cfg = yaml.safe_load((root / "config.yaml").read_text())
    print(write_request(root / "creator_request_template.xlsx", cfg["platforms"]))
